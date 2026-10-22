import traceback
from django.db.models import Q, Min
from django.db.models.functions import Cast
from django.db.models import DateTimeField, IntegerField
from django.db import transaction
from django.http import HttpResponse
from django.http import FileResponse
from django.conf import settings
from django.contrib import messages
from django.views.decorators.csrf import csrf_exempt
from django.utils import timezone
from django.contrib.auth.models import User
from django.contrib.gis.geos import GEOSGeometry
from django.contrib.gis.db.models.functions import Transform
from django.core.exceptions import ValidationError

#from rest_framework import viewsets, permissions

from rest_framework.views import APIView
from rest_framework import viewsets, serializers, status, generics, views
#from rest_framework.decorators import detail_route, list_route,renderer_classes
from rest_framework.decorators import action
from rest_framework.decorators import action as detail_route
from rest_framework.decorators import action as list_route
from rest_framework.decorators import renderer_classes
from rest_framework.response import Response
from rest_framework.renderers import JSONRenderer
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from rest_framework.permissions import IsAuthenticated, AllowAny, IsAdminUser, BasePermission, SAFE_METHODS
from rest_framework.pagination import PageNumberPagination
from rest_framework_datatables.pagination import DatatablesPageNumberPagination
from rest_framework_datatables.filters import DatatablesFilterBackend

from datetime import datetime, timedelta, date
import json
import os

from silrec.helpers import is_customer, is_internal
from silrec.components.proposals.models import   (
    TextSearchFieldDisplay,
    Proposal,
)
from silrec.components.lookups.models import (
    ObjectiveLkp,
    RegenerationMethodsLkp,
    SpeciesApiLkp,
    TaskLkp,
)
from silrec.components.forest_blocks.models import   (
    Polygon,
    #TmpPolygon,
    Cohort,
    Treatment,
    Operation,
    TreatmentXtra,
    AssignChtToPly,
    Prescription,
    SilviculturistComment,
    SurveyAssessmentDocument,
)
from silrec.components.users.serializers import   (
    UserSerializer,
    UserSerializerSimple,
)
from silrec.components.forest_blocks.serializers import   (
    SurveyAssessmentDocumentSerializer,
    TreatmentSerializer,
    TreatmentXtraSerializer,
    CohortSerializer,
    PrescriptionSerializer,
    SilviculturistCommentSerializer,
    OperationSerializer,
    SimpleCohortSerializer,
    PolygonSerializer,
    Polygon2Serializer,
    PolygonSearchSerializer,
    PolygonCohortSerializer,
    PolygonGeometrySerializer,
    PolygonCohortDataSerializer,
)

import logging
logger = logging.getLogger(__name__)

#from .models import Cohort, Treatment, TreatmentXtra
#from .serializers import CohortSerializer, TreatmentSerializer, TreatmentXtraSerializer

class GetProfile(views.APIView):
    renderer_classes = [JSONRenderer,]
    def get(self, request, format=None):
        serializer  = UserSerializer(request.user,
                context={'request': request}
                )
        return Response(serializer.data)


class SurveyAssessmentDocumentViewSet(viewsets.ModelViewSet):
    queryset = SurveyAssessmentDocument.objects.all()
    serializer_class = SurveyAssessmentDocumentSerializer
    pagination_class = DatatablesPageNumberPagination
    filter_backends = (DatatablesFilterBackend,)

    def get_queryset(self):
        queryset = SurveyAssessmentDocument.objects.all().select_related('treatment')

        # Filter by treatment_id if provided
        treatment_id = self.request.query_params.get('treatment_id')
        if treatment_id:
            queryset = queryset.filter(treatment_id=treatment_id)

        # Apply additional filters
        document_type = self.request.query_params.get('document_type')
        status = self.request.query_params.get('status')
        date_from = self.request.query_params.get('date_from')
        date_to = self.request.query_params.get('date_to')
        search = self.request.query_params.get('search')

        if document_type and document_type != 'all':
            queryset = queryset.filter(document_type=document_type)

        if status and status != 'all':
            queryset = queryset.filter(status=status)

        if date_from:
            try:
                date_from = datetime.strptime(date_from, '%Y-%m-%d').date()
                queryset = queryset.filter(document_date__gte=date_from)
            except ValueError:
                pass

        if date_to:
            try:
                date_to = datetime.strptime(date_to, '%Y-%m-%d').date()
                queryset = queryset.filter(document_date__lte=date_to)
            except ValueError:
                pass

        if search:
            queryset = queryset.filter(
                Q(title__icontains=search) |
                Q(description__icontains=search) |
                Q(file_name__icontains=search)
            )

        return queryset.order_by('-document_date', '-created_on')

    def get_permissions(self):
        if self.action in ['list', 'retrieve', 'create', 'update', 'partial_update']:
            permission_classes = [IsAuthenticated]
        else:
            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
        return [permission() for permission in permission_classes]

#    @action(detail=True, methods=['get'])
#    def download(self, request, pk=None):
#        """Download document file"""
#        document = self.get_object()
#
#        if document.file:
#            response = HttpResponse(document.file, content_type='application/octet-stream')
#            response['Content-Disposition'] = f'attachment; filename="{document.file_name}"'
#            return response
#        else:
#            return Response({'error': 'No file attached'}, status=404)

    @action(detail=True, methods=['get'])
    def download(self, request, pk=None):
        """Download document file"""
        document = self.get_object()

        if document.file:
            # Get the original filename from the database
            filename = document.file_name or 'document'

            # Open the file
            file_handle = document.file.open()

            # Create response with correct filename
            response = FileResponse(file_handle, content_type='application/octet-stream')
            response['Content-Disposition'] = f'attachment; filename="{filename}"'
            return response
        else:
            return Response({'error': 'No file attached'}, status=404)

    @action(detail=False, methods=['get'])
    def document_types(self, request):
        """Get available document types"""
        types = [{'value': value, 'label': label} for value, label in SurveyAssessmentDocument.DOCUMENT_TYPES]
        return Response(types)

    @action(detail=False, methods=['get'])
    def status_choices(self, request):
        """Get available status choices"""
        statuses = [{'value': value, 'label': label} for value, label in SurveyAssessmentDocument.STATUS_CHOICES]
        return Response(statuses)

    # Add this method to handle partial updates (including marked_deleted)
    def partial_update(self, request, *args, **kwargs):
        """Handle partial updates including marked_deleted field"""
        #marked_deleted = request.data.get('marked_deleted')
        instance = self.get_object()
        #instance.marked_deleted = marked_deleted
        serializer = self.get_serializer(instance, data=request.data, partial=True)
        #import ipdb; ipdb.set_trace()
        serializer.is_valid(raise_exception=True)
        self.perform_update(serializer)

        return Response(serializer.data)

    def perform_create(self, serializer):
        """Set uploaded_by when creating a document"""
        #import ipdb; ipdb.set_trace()
        serializer.save(uploaded_by=self.request.user)

#    def perform_update(self, serializer):
#        """Perform update, handling file changes if needed"""
#        # Check if a new file is being uploaded
#        file = self.request.FILES.get('file')
#        if file:
#            # If a new file is uploaded, update file metadata
#            instance = serializer.save()
#            instance.file_name = file.name
#            instance.file_size = file.size
#            instance.save()
#        else:
#            instance = self.get_object()
#            filename = os.path.basename(instance.file_name)
#            serializer.save()
#            instance.file_name = filename
#            instance.save()
#            import ipdb; ipdb.set_trace()
#            pass


class PolygonCohortViewSet(viewsets.ModelViewSet):
    queryset = AssignChtToPly.objects.none()
    serializer_class = PolygonCohortSerializer

    def get_queryset(self):
        #import ipdb; ipdb.set_trace()
        return AssignChtToPly.objects.all() #[:5]

    def list(self, request, *args, **kwargs):
        """ http://localhost:8001/api/polygoncohorts.json
        """
        serializer = PolygonCohortSerializer(self.get_queryset()[:5], many=True)
        return Response(serializer.data)

class PolygonViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = PolygonSerializer

    def get_queryset(self):
        #import ipdb; ipdb.set_trace()
        return Polygon.objects.all() #[:5]

    def list(self, request, *args, **kwargs):
        """ http://localhost:8001/api/polygon.json
        """
        serializer = PolygonSerializer(self.get_queryset()[:5], many=True)
        return Response(serializer.data)


class Polygon2ViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = Polygon2Serializer

    def get_queryset(self):
        #import ipdb; ipdb.set_trace()
        return Polygon.objects.all() #[:5]

#    def list(self, request, *args, **kwargs):
#        """ http://localhost:8001/api/polygon.json
#        """
#        serializer = Polygon2Serializer(self.get_queryset()[:5], many=True)
#        import ipdb; ipdb.set_trace()
#
#        return Response(serializer.data)


class PolygonGeometryViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = PolygonGeometrySerializer

    def get_queryset(self):
        #import ipdb; ipdb.set_trace()
        return Polygon.objects.all() #[:5]

#    def list(self, request, *args, **kwargs):
#        """ http://localhost:8001/api/polygon.json
#        """
#        serializer = PolygonGeometrySerializer(self.get_queryset()[:5], many=True)
#        import ipdb; ipdb.set_trace()
#
#        return Response(serializer.data)



class PolygonPaginatedViewSet(viewsets.ModelViewSet):
    filter_backends = (DatatablesFilterBackend,)
    #filter_backends = (ProposalFilterBackend,)
    pagination_class = DatatablesPageNumberPagination
    #renderer_classes = (ProposalRenderer,)
    #queryset = Polygon.objects.none()
    queryset = Polygon.objects.filter(polygon_id__lt=392915)
    #serializer_class = ListProposalSerializer
    serializer_class = Polygon2Serializer
    search_fields = ['polygon_id',]
    #serializer_class = DTProposalSerializer
    page_size = 10

    def get_queryset(self):
        #import ipdb; ipdb.set_trace()
        #return Polygon.objects.all() #[:5]
        #return Polygon.objects.all()[:25]
        return Polygon.objects.filter(polygon_id__lt=392915)

    #@list_route(methods=['GET', ])
#    @action(detail=False, methods=['GET'])
#    @list_route(
#        methods=[
#            "GET",
#        ],
#        detail=False,
#    )
    @list_route(detail=False, methods=['GET', ])
    def ply_datatable_list(self, request, *args, **kwargs):
        """ http://localhost:8001/api/ply_paginated/ply_datatable_list/?format=datatables&draw=1&length=10
        """
        #self.serializer_class = DTSpatialQueryLayersUsedSerializer
        self.serializer_class = Polygon2Serializer
        #queryset = self.get_queryset().filter(layer_data__isnull=False, processing_status=Proposal.PROCESSING_STATUS_APPROVED)
        queryset = self.get_queryset().filter()

        queryset = self.filter_queryset(queryset)
        self.paginator.page_size = queryset.count()
        # self.paginator.page_size = 0
        result_page = self.paginator.paginate_queryset(queryset, request)
        serializer = Polygon2Serializer(
            result_page, context={'request': request}, many=True
        )
        data = serializer.data

        response = self.paginator.get_paginated_response(data)
        return response

    #@list_route(methods=['GET',])
    @list_route(detail=False, methods=['GET', ])
    def list_paginated(self, request, *args, **kwargs):
        """
        http://localhost:8001/api/ply_paginated/list_paginated/?format=datatables&draw=1&length=10
        """
        #import ipdb; ipdb.set_trace()
        proposals = self.get_queryset()
        paginator = PageNumberPagination()
        #paginator = LimitOffsetPagination()
        paginator.page_size = 2
        result_page = paginator.paginate_queryset(proposals, request)
        serializer = Polygon2Serializer(result_page, context={'request':request}, many=True)
        return paginator.get_paginated_response(serializer.data)


from silrec.components.forest_blocks.models import Polygon, AssignChtToPly, Cohort
class DebugPolygonRelationsView(views.APIView):
    ''' Check relation names
        http://localhost:8001/api/debug-polygon-relations/
    '''

    def get(self, request):
        from django.db.models.fields.related import ForeignObjectRel
        from silrec.components.forest_blocks.models import Polygon

        polygon = Polygon.objects.first()
        if polygon:
            # Get all related fields
            related_objects = [
                f for f in polygon._meta.get_fields()
                if (f.one_to_many or f.one_to_one) and f.auto_created
            ]

            relations_info = []
            for rel in related_objects:
                relations_info.append({
                    'name': rel.get_accessor_name(),
                    'related_model': rel.related_model.__name__,
                    'field_name': rel.field.name if hasattr(rel, 'field') else 'N/A'
                })

            return Response({
                'polygon_id': polygon.polygon_id,
                'related_fields': relations_info
            })
        return Response({'error': 'No Polygon objects found'})


class DatatablesPageNumberPagination(PageNumberPagination):
    page_size_query_param = 'length'
    page_query_param = 'page'
    max_page_size = 1000
    page_size = 10

    def get_page_number(self, request, paginator):
        # Calculate page number from DataTables 'start' parameter
        start = request.query_params.get('start')
        length = request.query_params.get('length', self.page_size)

        if start and length:
            try:
                start = int(start)
                length = int(length)
                if length > 0:
                    page_number = (start // length) + 1
                    return max(1, page_number)
            except (ValueError, TypeError):
                pass

        # Fallback to default page number
        return super().get_page_number(request, paginator)

    def get_paginated_response(self, data):
        # Get datatables parameters
        draw = int(self.request.query_params.get('draw', 1))

        # Get counts
        total_count = getattr(self.request, '_datatables_total_count', self.page.paginator.count)
        filtered_count = getattr(self.request, '_datatables_filtered_count', self.page.paginator.count)

        return Response({
            'draw': draw,
            'recordsTotal': total_count,
            'recordsFiltered': filtered_count,
            'data': data
        })


class DatatablesFilterBackend(DatatablesFilterBackend):
    """
    Custom filter backend for datatables
    """
    def filter_queryset(self, request, queryset, view):
        total_count = queryset.count()

        # Apply your custom filters here
        filter_polygon_name = request.query_params.get('filter_polygon_name', '')
        filter_species = request.query_params.get('filter_species', '')
        filter_min_area = request.query_params.get('filter_min_area', '')
        filter_cohort_status = request.query_params.get('filter_cohort_status', 'all')

        if filter_polygon_name:
            queryset = queryset.filter(name__icontains=filter_polygon_name)

        if filter_min_area:
            try:
                min_area = float(filter_min_area)
                queryset = queryset.filter(area_ha__gte=min_area)
            except ValueError:
                pass

        if filter_species:
            queryset = queryset.filter(
                assignchttoply__cohort__species__icontains=filter_species,
                assignchttoply__status_current=True
            ).distinct()

        if filter_cohort_status != 'all':
            status_active = filter_cohort_status == 'active'
            queryset = queryset.filter(
                assignchttoply__status_current=status_active
            ).distinct()

        filtered_count = queryset.count()

        setattr(view, '_datatables_total_count', total_count)
        setattr(view, '_datatables_filtered_count', filtered_count)

        return queryset


class PolygonCohortTableViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = PolygonCohortDataSerializer
    pagination_class = DatatablesPageNumberPagination
    filter_backends = (DatatablesFilterBackend,)

    # DataTables column name -> ORM ordering expression.
    ORDERING_FIELDS = {
        'polygon_id': 'polygon_id',
        'name': 'name',
        'zfea_id': 'zfea_id',
        'area_ha': 'area_ha',
        'cohort_id': 'assignchttoply__cohort__cohort_id',
        'obj_code': 'assignchttoply__cohort__obj_code',
        'target_ba_m2ha': 'assignchttoply__cohort__target_ba_m2ha',
        'resid_ba_m2ha': 'assignchttoply__cohort__resid_ba_m2ha',
        'species': 'assignchttoply__cohort__species',
        'status_current': 'assignchttoply__status_current',
    }

    # Multi-valued relations are aggregated with Min() so a polygon is not
    # duplicated across rows when sorted on these columns.
    AGGREGATED_ORDERING_FIELDS = {
        'assignchttoply__cohort__cohort_id',
        'assignchttoply__cohort__obj_code',
        'assignchttoply__cohort__target_ba_m2ha',
        'assignchttoply__cohort__resid_ba_m2ha',
        'assignchttoply__cohort__species',
        'assignchttoply__status_current',
    }

    # Boolean relations have no MIN() in Postgres — cast to integer first.
    BOOLEAN_ORDERING_FIELDS = {
        'assignchttoply__status_current',
    }

    def get_queryset(self):
        proposal_id = self.request.query_params.get('proposal_id')
        if not proposal_id:
            return Polygon.objects.none()

        try:
            # Keep it simple - use only select_related
            return Polygon.objects.filter(
                proposal_id=proposal_id
            ).select_related(
                'compartment',
                'sp_code'
            ).distinct()
        except Exception as e:
            logger.error(f"Error in get_queryset: {str(e)}")
            return Polygon.objects.none()

    def get_ordering(self, request):
        """Resolve DataTables order[i][column] params against column names."""
        columns = {}
        i = 0
        while True:
            name = request.query_params.get('columns[%d][name]' % i)
            data = request.query_params.get('columns[%d][data]' % i)
            if name is None and data is None:
                break
            columns[i] = name or data
            i += 1

        ordering = []
        i = 0
        while True:
            idx = request.query_params.get('order[%d][column]' % i)
            if idx is None:
                break
            direction = request.query_params.get('order[%d][dir]' % i, 'asc')
            field_name = columns.get(int(idx))
            field = self.ORDERING_FIELDS.get(field_name)
            if field:
                ordering.append(('-' if direction == 'desc' else '') + field)
            i += 1

        return ordering

    def apply_ordering(self, queryset, request):
        """Apply DataTables ordering, aggregating multi-valued relations."""
        ordering = self.get_ordering(request)
        if not ordering:
            return queryset

        annotations = {}
        order_by = []
        for field in ordering:
            desc = field.startswith('-')
            name = field[1:] if desc else field
            if name in self.AGGREGATED_ORDERING_FIELDS:
                annotation = 'order_%s' % name.replace('__', '_')
                if name in self.BOOLEAN_ORDERING_FIELDS:
                    annotations[annotation] = Min(
                        Cast(name, output_field=IntegerField())
                    )
                else:
                    annotations[annotation] = Min(name)
                order_by.append(('-' if desc else '') + annotation)
            else:
                order_by.append(field)

        if annotations:
            queryset = queryset.annotate(**annotations)

        return queryset.order_by(*order_by)

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        queryset = self.apply_ordering(queryset, request)

        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)


class _PolygonCohortTableViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = PolygonCohortDataSerializer
    pagination_class = DatatablesPageNumberPagination
    filter_backends = (DatatablesFilterBackend,)  # Use custom filter backend
    #page_size = 2

    def get_queryset(self):
        proposal_id = self.request.query_params.get('proposal_id')
        if proposal_id:
            return Polygon.objects.filter(
                proposal_id=proposal_id
            ).prefetch_related(
                #'tmpassignchttoply_set',
                #'tmpassignchttoply_set__cohort'
                'assignchttoply_set',
                'assignchttoply_set__cohort'
            )
        return Polygon.objects.none()

#    def _list(self, request, *args, **kwargs):
#        """
#        http://localhost:8001/api/ply_paginated/list_paginated/?format=datatables&draw=1&length=10
#        http://localhost:8001/api/polygon_cohort_table/?fmt=datatables&draw=3&length=10
#        """
#        #import ipdb; ipdb.set_trace()
#        polygons = self.get_queryset()
#        paginator = PageNumberPagination()
#        #paginator = LimitOffsetPagination()
#        paginator.page_size = 2
#        result_page = paginator.paginate_queryset(polygons, request)
#        serializer = PolygonCohortDataSerializer(result_page, context={'request':request}, many=True)
#        return paginator.get_paginated_response(serializer.data)

class IsOfficer(BasePermission):
    def has_permission(self, request, view):
        return True

class IsOperator(BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.groups.filter(name='Operator').exists()

class IsReviewer(BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and request.user.groups.filter(name='Reviewer').exists()

class IsSilrecAdmin(BasePermission):
    def has_permission(self, request, view):
        return request.user.is_authenticated and (
            request.user.groups.filter(name='Silrec Admin').exists() or
            request.user.is_superuser
        )

class ReadOnlyPermission(BasePermission):
    def has_permission(self, request, view):
        return True
#        if request.method in SAFE_METHODS:
#            return True
#        return False

class CohortViewSet(viewsets.ModelViewSet):
    queryset = Cohort.objects.all()
    serializer_class = CohortSerializer

#    def get_queryset(self):
#        #import ipdb; ipdb.set_trace()
#        print(self.request.query_params)
#        pass

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            permission_classes = [IsAuthenticated]
        else:
            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
        return [permission() for permission in permission_classes]


class TreatmentDatatablesFilterBackend(DatatablesFilterBackend):
    """
    Custom filter backend for datatables treatment filtering
    """
    def filter_queryset(self, request, queryset, view):
        # Get total count before filtering
        total_count = queryset.count()

        # Apply custom filters for treatments
        filter_task = request.query_params.get('filter_task', '')
        filter_task_classification_id = request.query_params.get('filter_task_classification_id', '')
        filter_status = request.query_params.get('filter_status', 'all')
        filter_plan_year = request.query_params.get('filter_plan_year', '')
        filter_plan_month = request.query_params.get('filter_plan_month', 'all')
        filter_complete_date_from = request.query_params.get('filter_complete_date_from', '')
        filter_complete_date_to = request.query_params.get('filter_complete_date_to', '')
        filter_machine = request.query_params.get('filter_machine', '')
        filter_operator = request.query_params.get('filter_operator', '')

        # Filter by task
        if filter_task:
            queryset = queryset.filter(task__task=filter_task)

        # Filter by task classification
        if (
            filter_task_classification_id
            and filter_task_classification_id != 'all'
        ):
            try:
                task_class_id = int(filter_task_classification_id)
                queryset = queryset.filter(
                    task__task_class_id=task_class_id
                )
            except ValueError:
                pass

        # Filter by status
        if filter_status != 'all':
            queryset = queryset.filter(status=filter_status)

        # Filter by planned year
        if filter_plan_year:
            try:
                plan_year = int(filter_plan_year)
                queryset = queryset.filter(plan_yr=plan_year)
            except ValueError:
                pass

        # Filter by planned month
        if filter_plan_month != 'all':
            try:
                plan_month = int(filter_plan_month)
                queryset = queryset.filter(plan_mth=plan_month)
            except ValueError:
                pass

        # Filter by complete date range
        if filter_complete_date_from:
            try:
                complete_date_from = datetime.strptime(filter_complete_date_from, '%Y-%m-%d').date()
                queryset = queryset.filter(complete_date__gte=complete_date_from)
            except ValueError:
                pass

        if filter_complete_date_to:
            try:
                complete_date_to = datetime.strptime(filter_complete_date_to, '%Y-%m-%d').date()
                queryset = queryset.filter(complete_date__lte=complete_date_to)
            except ValueError:
                pass

        # Filter by machine (from TreatmentXtra)
        if filter_machine:
            queryset = queryset.filter(
                treatmentxtra__zmachine_id__icontains=filter_machine
            ).distinct()

        # Filter by operator
        if filter_operator:
            queryset = queryset.filter(
                changed_by__icontains=filter_operator
            )

        # Filter by Post 2024 only — treatments whose cohort links to an
        # Operation with plan_period in OP_PLAN_PERIOD.
        filter_post_2024_only = request.query_params.get('filter_post_2024_only', '')
        if filter_post_2024_only in ('true', 'True', True):
            plan_periods = settings.OP_PLAN_PERIOD
            queryset = queryset.filter(
                cohort__op_id__in=Operation.objects.filter(
                    plan_period__in=plan_periods
                ).values('op_id')
            )

        filtered_count = queryset.count()

        # Set counts on the request for the paginator to use
        setattr(request, '_datatables_total_count', total_count)
        setattr(request, '_datatables_filtered_count', filtered_count)

        return queryset


class TreatmentViewSet(viewsets.ModelViewSet):
    queryset = Treatment.objects.all()
    serializer_class = TreatmentSerializer
    pagination_class = DatatablesPageNumberPagination
    filter_backends = (TreatmentDatatablesFilterBackend,)

    def get_queryset(self):
        # Start with a properly ordered queryset
        queryset = Treatment.objects.all().select_related('task', 'cohort').order_by('treatment_id')

        # Handle cohort filtering
        cohort_id = self.request.query_params.get('cohort_id')
        if cohort_id:
            queryset = queryset.filter(cohort_id=cohort_id)

        return queryset

    def list(self, request, *args, **kwargs):
        # Always use datatables format for list requests
        return self.datatables_list(request, *args, **kwargs)

    def datatables_list(self, request, *args, **kwargs):
        """Custom list method for datatables format"""
        queryset = self.filter_queryset(self.get_queryset())

        # Apply ordering from DataTables if provided
        ordering = self.get_ordering(request, queryset)
        if ordering:
            queryset = queryset.order_by(*ordering)

        # Apply pagination
        page = self.paginate_queryset(queryset)
        if page is not None:
            serializer = self.get_serializer(page, many=True)
            return self.get_paginated_response(serializer.data)

        serializer = self.get_serializer(queryset, many=True)
        return Response(serializer.data)

    def get_ordering(self, request, queryset):
        """Extract ordering from DataTables request"""
        order_column_index = request.query_params.get('order[0][column]')
        order_direction = request.query_params.get('order[0][dir]', 'asc')

        if order_column_index is not None:
            # Map DataTables column index to model field
            column_map = {
                '0': 'treatment_id',  # ID column
                '1': 'task__name',     # Task column
                '2': 'plan_yr',        # Planned Year column
                '3': 'plan_mth',       # Planned Month column
                '4': 'status',         # Status column
                '5': 'complete_date',  # Complete Date column
            }

            field_name = column_map.get(order_column_index)
            if field_name:
                if order_direction == 'desc':
                    field_name = '-' + field_name
                return [field_name]

        # Default ordering if none specified
        return ['treatment_id']

    def get_permissions(self):
        print(f"TreatmentViewSet - Action: {self.action}")
        print(f"TreatmentViewSet - User: {self.request.user}")

        if self.action in ['list', 'retrieve']:
            permission_classes = [IsAuthenticated]
        else:
            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]

        return [permission() for permission in permission_classes]


class TreatmentXtraViewSet(viewsets.ModelViewSet):
    queryset = TreatmentXtra.objects.all()
    serializer_class = TreatmentXtraSerializer

#    def get_permissions(self):
#        import ipdb; ipdb.set_trace()
#        if self.action in ['list', 'retrieve', 'update', 'partial_update']:
#            permission_classes = [IsAuthenticated]
#        else:
#            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
#        return [permission() for permission in permission_classes]

    def get_queryset(self):
        queryset = TreatmentXtra.objects.all()
        treatment_id = self.request.query_params.get('treatment_id')
        if treatment_id:
            queryset = queryset.filter(treatment_id=treatment_id)
        return queryset


class PrescriptionViewSet(viewsets.ModelViewSet):
    queryset = Prescription.objects.all()
    serializer_class = PrescriptionSerializer

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            permission_classes = [IsAuthenticated]
        else:
            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
        return [permission() for permission in permission_classes]

    def get_queryset(self):
        queryset = Prescription.objects.all()
        #prescription_id = self.request.query_params.get('prescription_id')
        prescription_id = self.request.query_params.get('prescription')
        if prescription_id:
            queryset = queryset.filter(prescription_id=prescription_id)
        return queryset


class SilviculturistCommentViewSet(viewsets.ModelViewSet):
    queryset = SilviculturistComment.objects.all()
    serializer_class = SilviculturistCommentSerializer

    def get_queryset(self):
        queryset = SilviculturistComment.objects.all()
        treatment_id = self.request.query_params.get('treatment')
        if treatment_id:
            queryset = queryset.filter(treatment_id=treatment_id)
        return queryset

    def get_permissions(self):
        if self.action in ['list', 'retrieve']:
            permission_classes = [IsAuthenticated]
        else:
            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
        return [permission() for permission in permission_classes]


class PolygonSearchViewSet(viewsets.ModelViewSet):
    queryset = Polygon.objects.none()
    serializer_class = PolygonSearchSerializer
    pagination_class = DatatablesPageNumberPagination
    filter_backends = (DatatablesFilterBackend,)

    def get_queryset(self):
        # Transform geometry to EPSG:4326 at the database level
        queryset = Polygon.objects.all().annotate(
            geom_4326=Transform('geom', 4326),
            created_on_dt=Cast('created_on', DateTimeField())
        ).prefetch_related(
            'compartment',
            'assignchttoply_set',
            'assignchttoply_set__op',
            'assignchttoply_set__cohort',
            'assignchttoply_set__cohort__treatment_set',
            'assignchttoply_set__cohort__treatment_set__task'
        ).select_related(
            'compartment',
            'sp_code'
        ).order_by('polygon_id')

        # Apply filters
        obj_code = self.request.query_params.get('obj_code')
        obj_classification_id = self.request.query_params.get('obj_classification_id')
        compartment = self.request.query_params.get('compartment')
        block = self.request.query_params.get('block')
        district = self.request.query_params.get('district')
        zfea_id = self.request.query_params.get('zfea_id')
        treatment_status = self.request.query_params.get('treatment_status')
        created_from = self.request.query_params.get('created_from')
        created_to = self.request.query_params.get('created_to')
        return_empty = self.request.query_params.get('return_empty')

        # Filter by objective classification
        if (
            obj_classification_id
            and obj_classification_id != 'all'
        ):
            try:
                class_id = int(obj_classification_id)
                obj_codes = ObjectiveLkp.objects.filter(
                    obj_class_id=class_id
                ).values('obj_code')
                queryset = queryset.filter(
                    assignchttoply__cohort__obj_code__in=obj_codes,
                    assignchttoply__status_current=True
                ).distinct()
            except ValueError:
                pass

        if obj_code:
            queryset = queryset.filter(
                assignchttoply__cohort__obj_code__iexact=obj_code.strip(),
                assignchttoply__status_current=True
            ).distinct()

        if compartment:
            queryset = queryset.filter(compartment__compartment__icontains=compartment)

        if block:
            queryset = queryset.filter(compartment__block__icontains=block)

        if district:
            queryset = queryset.filter(compartment__district__icontains=district)

        # Filter by post 2024 only — polygons whose current cohort links to an
        # Operation with plan_period in OP_PLAN_PERIOD.
        filter_post_2024_only = self.request.query_params.get('filter_post_2024_only')
        if filter_post_2024_only and filter_post_2024_only.lower() == 'true':
            plan_periods = settings.OP_PLAN_PERIOD
            queryset = queryset.filter(
                assignchttoply__cohort__op_id__in=Operation.objects.filter(
                    plan_period__in=plan_periods
                ).values('op_id'),
                assignchttoply__status_current=True
            ).distinct()

        if zfea_id:
            # FEA ID is carried on the current assignment's operation; the
            # polygon.zfea_id column is legacy/unpopulated. Fall back to it too.
            queryset = queryset.filter(
                Q(assignchttoply__op__fea_id__icontains=zfea_id) |
                Q(zfea_id__icontains=zfea_id)
            ).distinct()

        if treatment_status:
            queryset = queryset.filter(
                assignchttoply__cohort__treatment__status=treatment_status,
                assignchttoply__status_current=True
            ).distinct()

        if created_from:
            try:
                created_from_date = datetime.strptime(created_from, '%Y-%m-%d').date()
                queryset = queryset.filter(created_on_dt__gte=created_from_date)
            except ValueError:
                pass

        if created_to:
            try:
                created_to_date = datetime.strptime(created_to, '%Y-%m-%d').date()
                queryset = queryset.filter(created_on_dt__lte=created_to_date)
            except ValueError:
                pass

        return queryset

    # Map DataTables column names (sent as columns[i][name]) to an ordering
    # expression. Multi-valued relations are aggregated with Min() so a polygon
    # is not duplicated across rows when sorted by e.g. obj_code.
    ORDERING_FIELDS = {
        'polygon_id': 'polygon_id',
        'name': 'name',
        'compartment__compartment': 'compartment__compartment',
        'area_ha': 'area_ha',
        'zfea_id': 'zfea_id',
        'fea_id': 'assignchttoply__op__fea_id',
        'assignchttoply__cohort__obj_code': 'assignchttoply__cohort__obj_code',
        'assignchttoply__cohort__species': 'assignchttoply__cohort__species',
        'assignchttoply__cohort__target_ba_m2ha': 'assignchttoply__cohort__target_ba_m2ha',
        'assignchttoply__cohort__resid_ba_m2ha': 'assignchttoply__cohort__resid_ba_m2ha',
        'assignchttoply__cohort__treatment__status': 'assignchttoply__cohort__treatment__status',
        'created_on': 'created_on',
    }
    AGGREGATED_ORDERING_FIELDS = {
        'assignchttoply__op__fea_id',
        'assignchttoply__cohort__obj_code',
        'assignchttoply__cohort__species',
        'assignchttoply__cohort__target_ba_m2ha',
        'assignchttoply__cohort__resid_ba_m2ha',
        'assignchttoply__cohort__treatment__status',
    }

    def get_ordering(self, request):
        """Build an ordering list from the DataTables order[i][column] params.

        DataTables sends the column index and direction; we resolve the index to
        the column name configured in the Vue component (columns[i][name]).
        """
        columns = {}
        i = 0
        while True:
            name = request.query_params.get('columns[%d][name]' % i)
            data = request.query_params.get('columns[%d][data]' % i)
            if name is None and data is None:
                break
            columns[i] = name or data
            i += 1

        ordering = []
        i = 0
        while True:
            idx = request.query_params.get('order[%d][column]' % i)
            if idx is None:
                break
            direction = request.query_params.get('order[%d][dir]' % i, 'asc')
            field_name = columns.get(int(idx))
            field = self.ORDERING_FIELDS.get(field_name)
            if field:
                ordering.append(('-' if direction == 'desc' else '') + field)
            i += 1

        return ordering

    def apply_ordering(self, queryset, request):
        """Apply DataTables ordering, aggregating multi-valued relations."""
        ordering = self.get_ordering(request)
        if not ordering:
            return queryset

        annotations = {}
        order_by = []
        for field in ordering:
            desc = field.startswith('-')
            name = field[1:] if desc else field
            if name in self.AGGREGATED_ORDERING_FIELDS:
                annotation = 'order_%s' % name.replace('__', '_')
                annotations[annotation] = Min(name)
                order_by.append(('-' if desc else '') + annotation)
            else:
                order_by.append(field)

        if annotations:
            queryset = queryset.annotate(**annotations)

        return queryset.order_by(*order_by)

    def list(self, request, *args, **kwargs):
        queryset = self.filter_queryset(self.get_queryset())
        queryset = self.apply_ordering(queryset, request)

        # Apply pagination
        page = self.paginate_queryset(queryset)
        if page is not None:
            # Manually serialize to handle geometry transformation
            serializer = self.get_serializer(page, many=True)

            # Transform geometry data in the response
            response_data = self.transform_geometry_in_response(serializer.data)

            return self.get_paginated_response(response_data)

        serializer = self.get_serializer(queryset, many=True)
        response_data = self.transform_geometry_in_response(serializer.data)
        return Response(response_data)

    def transform_geometry_in_response(self, data):
        """
        Transform geometry data in the response to ensure it's in EPSG:4326
        """
        for item in data:
            if 'geom' in item and item['geom']:
                try:
                    # If we have the transformed geometry from the annotation, use it
                    polygon = Polygon.objects.filter(polygon_id=item['polygon_id']).annotate(
                        geom_4326=Transform('geom', 4326)
                    ).first()

                    if polygon and hasattr(polygon, 'geom_4326'):
                        # Convert to GeoJSON format
                        item['geom'] = json.loads(polygon.geom_4326.geojson)
                    else:
                        # Fallback: transform using GEOSGeometry
                        geom = GEOSGeometry(item['geom'])
                        geom.transform(4326)
                        item['geom'] = json.loads(geom.geojson)

                except Exception as e:
                    print(f"Error transforming geometry for polygon {item['polygon_id']}: {e}")
                    # Keep original geometry if transformation fails
                    pass

        return data


class OperationViewSet(viewsets.ModelViewSet):
    queryset = Operation.objects.all()
    serializer_class = OperationSerializer
    parser_classes = [MultiPartParser, FormParser, JSONParser]  # To handle file uploads
    pagination_class = DatatablesPageNumberPagination
    filter_backends = [DatatablesFilterBackend]

#    def get_permissions(self):
#        if self.action in ['list', 'retrieve', 'create', 'update', 'partial_update']:
#            permission_classes = [IsAuthenticated]
#        else:
#            permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]
#        return [permission() for permission in permission_classes]

    def get_queryset(self):
        queryset = Operation.objects.all()

        # Filter by cohort_id if provided
        cohort_id = self.request.query_params.get('cohort_id')
        if cohort_id:
            try:
                cohort = Cohort.objects.get(cohort_id=cohort_id)
                if cohort.op_id:
                    queryset = queryset.filter(op_id=cohort.op_id)
                else:
                    queryset = Operation.objects.none()
            except Cohort.DoesNotExist:
                queryset = Operation.objects.none()

        return queryset

    def perform_update(self, serializer):
        # BinaryField is read-only in DRF — handle clearing directly
        for field in ('silvic_plan_map', 'silvic_plan_doc'):
            val = self.request.data.get(field)
            if val is None or val == '':
                setattr(serializer.instance, field, None)
        serializer.save()

    @action(detail=False, methods=['get'])
    def by_cohort(self, request):
        """Get operation by cohort ID"""
        cohort_id = request.query_params.get('cohort_id')
        if not cohort_id:
            return Response(
                {'error': 'cohort_id parameter is required'},
                status=status.HTTP_400_BAD_REQUEST
            )

        try:
            cohort = Cohort.objects.get(cohort_id=cohort_id)
            if cohort.op_id:
                operation = Operation.objects.get(op_id=cohort.op_id)
                serializer = self.get_serializer(operation)
                return Response(serializer.data)
            else:
                # Check if there's a related polygon to get fea_id
                polygon = Polygon.objects.filter(
                    assignchttoply__cohort=cohort
                ).first()
                fea_id = polygon.zfea_id if polygon else None

                return Response({
                    'exists': False,
                    'fea_id': fea_id
                })
        except Cohort.DoesNotExist:
            return Response(
                {'error': 'Cohort not found'},
                status=status.HTTP_404_NOT_FOUND
            )
        except Operation.DoesNotExist:
            return Response(
                {'error': 'Operation not found'},
                status=status.HTTP_404_NOT_FOUND
            )

    def perform_create(self, serializer):
        """Handle operation creation and update cohort's op_id"""
        # Get cohort_id from request data or query params
        #import ipdb; ipdb.set_trace()
        cohort_id = self.request.data.get('cohort_id') or self.request.query_params.get('cohort_id')

        if cohort_id:
            try:
                cohort = Cohort.objects.get(cohort_id=cohort_id)
                # Save operation
                operation = serializer.save()
                # Update cohort's op_id
                cohort.op_id = operation.op_id
                cohort.save()
            except Cohort.DoesNotExist:
                # Still create the operation even if cohort doesn't exist
                serializer.save()
        else:
            serializer.save()

#    def partial_update(self, request, *args, **kwargs):
#        import ipdb; ipdb.set_trace()
#        pass
#
#    def perform_update(self, serializer):
#        import ipdb; ipdb.set_trace()
#        pass


class TestOperationUpdate(views.APIView):
    """Test endpoint to debug operation updates"""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        print("=== Test Operation Update ===")
        print("Request data:", request.data)
        print("Request FILES:", request.FILES)
        print("User:", request.user)

        # Try to get the operation
        op_id = request.data.get('op_id')
        if op_id:
            try:
                operation = Operation.objects.get(op_id=op_id)
                print("Operation found:", operation.op_id)

                # Try to update with serializer
                serializer = OperationSerializer(operation, data=request.data, partial=True)
                if serializer.is_valid():
                    print("Serializer is valid")
                    serializer.save()
                    return Response(serializer.data)
                else:
                    print("Serializer errors:", serializer.errors)
                    return Response(serializer.errors, status=400)
            except Operation.DoesNotExist:
                print("Operation not found")
                return Response({"error": "Operation not found"}, status=404)

        return Response({"error": "No operation ID provided"}, status=400)


# ---------------------------------------------------------------------------
# Bulk update across a set of polygons (cohort / treatments / operation).
#
# Two entry points:
#   POST /api/bulk-update/preview/  -> resolve polygon ids, return what would
#                                      be touched (no writes).
#   POST /api/bulk-update/          -> apply the changes in one transaction.
# ---------------------------------------------------------------------------
class _BulkUpdateBase(views.APIView):
    permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]

    # Whitelisted, user-settable fields per entity. Everything else in the
    # request is dropped so a client cannot reach unrelated columns.
    COHORT_FIELDS = {
        'obj_code', 'species', 'regen_method', 'site_quality',
        'target_ba_m2ha', 'resid_ba_m2ha', 'target_spha', 'resid_spha',
        'comments', 'op_date', 'pct_area', 'year_last_cut',
        'regen_date', 'regen_date2', 'regen_done', 'complete_date',
        'herbicide_app_spec', 'vrp', 'vrp_tot_area', 'stand',
        'extra_info', 'treatments',
    }
    TREATMENT_FIELDS = {
        'task', 'status', 'plan_yr', 'plan_mth', 'pct_area',
        'complete_date', 'organisation', 'results', 'reference',
        'initial_plan_yr', 'can_reschedule',
    }
    OPERATION_FIELDS = {
        'fea_id', 'das_id', 'plan_release', 'plan_period', 'document_date',
    }

    def _parse_polygon_ids(self, raw):
        """Accept a list of ints or a comma/space/newline/semicolon string."""
        if raw is None:
            return None, 'No polygon ids provided'
        if isinstance(raw, str):
            tokens = raw.replace(',', ' ').replace(';', ' ').split()
        elif isinstance(raw, (list, tuple)):
            tokens = [str(x).strip() for x in raw]
        else:
            return None, 'polygon_ids must be a list or a delimited string'

        ids, seen = [], set()
        for token in tokens:
            if not token:
                continue
            if not token.isdigit():
                return None, f"Invalid polygon id: '{token}'"
            n = int(token)
            if n not in seen:
                seen.add(n)
                ids.append(n)

        if not ids:
            return None, 'No polygon ids provided'
        return ids, None

    def _resolve(self, ids, require_proposal=True):
        """Resolve ids to polygons that have a current cohort.

        By default polygons must be linked to a proposal (the proposals
        dashboard drill-down). The Excel-upload path passes
        require_proposal=False so unlinked polygons are also accepted.
        Polygons without a current cohort (or missing) are reported as
        excluded with a reason.
        """
        polygons, excluded = [], []
        for pid in ids:
            poly = (Polygon.objects.select_related('compartment', 'proposal')
                    .filter(polygon_id=pid).first())
            if poly is None:
                excluded.append({'polygon_id': pid, 'name': None,
                                 'reason': 'Polygon not found'})
                continue
            if require_proposal and poly.proposal_id is None:
                excluded.append({'polygon_id': pid, 'name': poly.name,
                                 'reason': 'Not linked to a proposal (not available in the proposals dashboard)'})
                continue
            assignment = (AssignChtToPly.objects
                          .filter(polygon=poly, status_current=True)
                          .select_related('cohort')
                          .order_by('-cht2ply_id').first())
            if assignment is None or assignment.cohort_id is None:
                excluded.append({'polygon_id': pid, 'name': poly.name,
                                 'reason': 'No current cohort assigned'})
                continue
            polygons.append({
                'polygon': poly,
                'cohort': assignment.cohort,
                'assignment': assignment,
            })
        return polygons, excluded

    def _pick(self, allowed, fields):
        """Keep only whitelisted keys with a non-empty value."""
        if not fields:
            return {}
        out = {}
        for key, value in fields.items():
            if key not in allowed:
                continue
            if value is None or value == '':
                continue
            out[key] = value
        return out

    def _operation_for_cohort(self, cohort):
        if not cohort.op_id:
            return None
        return Operation.objects.filter(op_id=cohort.op_id).first()

    def _serializer_errors(self, serializer):
        errors = []
        for field, msgs in serializer.errors.items():
            errors.append(f"{field}: {', '.join(str(m) for m in msgs)}")
        return errors


class BulkUpdatePreviewView(_BulkUpdateBase):
    def post(self, request):
        ids, err = self._parse_polygon_ids(request.data.get('polygon_ids'))
        if err:
            return Response({'error': err}, status=status.HTTP_400_BAD_REQUEST)

        require_proposal = request.data.get('require_proposal', True)
        polygons, excluded = self._resolve(ids, require_proposal=require_proposal)
        rows = []
        for item in polygons:
            poly, cohort = item['polygon'], item['cohort']
            op = self._operation_for_cohort(cohort)
            rows.append({
                'polygon_id': poly.polygon_id,
                'name': poly.name,
                'compartment': poly.compartment.compartment if poly.compartment else None,
                'proposal_id': poly.proposal_id,
                'cohort_id': cohort.cohort_id,
                'obj_code': cohort.obj_code,
                'species': cohort.species,
                'treatment_count': Treatment.objects.filter(cohort=cohort).count(),
                'operation': {'op_id': op.op_id, 'fea_id': op.fea_id} if op else None,
            })

        return Response({
            'requested': len(ids),
            'found': len(rows),
            'excluded': excluded,
            'polygons': rows,
        })


class BulkUpdateApplyView(_BulkUpdateBase):
    def post(self, request):
        payload = request.data or {}
        ids, err = self._parse_polygon_ids(payload.get('polygon_ids'))
        if err:
            return Response({'error': err}, status=status.HTTP_400_BAD_REQUEST)

        require_proposal = payload.get('require_proposal', True)
        polygons, excluded = self._resolve(ids, require_proposal=require_proposal)

        cohort_section = payload.get('cohort') or {}
        treatment_section = payload.get('treatments') or {}
        operation_section = payload.get('operation') or {}

        # Upfront validation (nothing written yet) so the user sees all errors.
        problems = self._validate_payload(request, cohort_section, treatment_section, operation_section)
        if problems:
            return Response({'error': 'Validation failed', 'details': problems},
                            status=status.HTTP_400_BAD_REQUEST)

        results = []
        try:
            with transaction.atomic():
                for item in polygons:
                    results.append(self._apply_one(
                        request, item, cohort_section, treatment_section, operation_section
                    ))
        except serializers.ValidationError as exc:
            return Response({'error': 'Validation failed', 'details': exc.detail},
                            status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            logger.exception('Bulk update failed and was rolled back')
            return Response(
                {'error': 'Bulk update failed and was rolled back. No changes were applied.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        return Response({
            'applied': len([r for r in results if not r.get('skipped')]),
            'excluded': excluded,
            'results': results,
        })

    def _validate_payload(self, request, cohort_section, treatment_section, operation_section):
        problems = []

        cohort_fields = self._pick(self.COHORT_FIELDS, cohort_section.get('fields'))
        if cohort_fields:
            ser = CohortSerializer(data=cohort_fields, partial=True, context={'request': request})
            if not ser.is_valid():
                problems.extend(self._serializer_errors(ser))

        treatment_mode = treatment_section.get('mode', 'create')
        treatment_fields = self._pick(self.TREATMENT_FIELDS, treatment_section.get('fields'))
        if treatment_section:
            if treatment_mode == 'create':
                task = (treatment_section.get('fields') or {}).get('task')
                if not task or not str(task).strip():
                    problems.append('treatments.task: Must select a Task')
                else:
                    ser = TreatmentSerializer(
                        data={**treatment_fields, 'task': task},
                        context={'request': request},
                    )
                    if not ser.is_valid():
                        problems.extend(self._serializer_errors(ser))
            elif treatment_mode == 'update':
                match = treatment_section.get('match') or {}
                if not (match.get('task') or match.get('status')):
                    problems.append('treatments.match: update mode requires a task and/or status filter')
                if treatment_fields:
                    ser = TreatmentSerializer(data=treatment_fields, partial=True, context={'request': request})
                    if not ser.is_valid():
                        problems.extend(self._serializer_errors(ser))
            else:
                problems.append(f"treatments.mode: unsupported mode '{treatment_mode}'")

        operation_mode = operation_section.get('mode', 'create')
        operation_fields = self._pick(self.OPERATION_FIELDS, operation_section.get('fields'))
        if operation_section:
            if operation_mode == 'create':
                fea_id = (operation_section.get('fields') or {}).get('fea_id')
                if not fea_id or not str(fea_id).strip():
                    problems.append('operation.fea_id: Must enter an FEA ID')
                else:
                    ser = OperationSerializer(data=operation_fields, context={'request': request})
                    if not ser.is_valid():
                        problems.extend(self._serializer_errors(ser))
            elif operation_mode == 'update':
                if operation_fields:
                    ser = OperationSerializer(data=operation_fields, partial=True, context={'request': request})
                    if not ser.is_valid():
                        problems.extend(self._serializer_errors(ser))
            else:
                problems.append(f"operation.mode: unsupported mode '{operation_mode}'")

        return problems

    def _apply_one(self, request, item, cohort_section, treatment_section, operation_section):
        poly, cohort, assignment = item['polygon'], item['cohort'], item['assignment']
        result = {
            'polygon_id': poly.polygon_id,
            'cohort_id': cohort.cohort_id,
            'cohort_updated': False,
            'treatments_created': [],
            'treatments_updated': [],
            'operation_created': None,
            'operation_updated': None,
            'notes': [],
            'skipped': False,
        }

        # Cohort: update the current cohort with the provided fields.
        cohort_fields = self._pick(self.COHORT_FIELDS, cohort_section.get('fields'))
        if cohort_fields:
            ser = CohortSerializer(instance=cohort, data=cohort_fields, partial=True,
                                   context={'request': request})
            ser.is_valid(raise_exception=True)
            ser.save()
            result['cohort_updated'] = True

        # Treatments.
        treatment_mode = treatment_section.get('mode', 'create')
        treatment_fields = self._pick(self.TREATMENT_FIELDS, treatment_section.get('fields'))
        if treatment_mode == 'create' and (treatment_fields or (treatment_section.get('fields') or {}).get('task')):
            ser = TreatmentSerializer(
                data={**treatment_fields, 'cohort': cohort.cohort_id},
                context={'request': request},
            )
            ser.is_valid(raise_exception=True)
            created = ser.save()
            result['treatments_created'].append(created.treatment_id)
        elif treatment_mode == 'update' and treatment_fields:
            qs = Treatment.objects.filter(cohort=cohort)
            match = treatment_section.get('match') or {}
            if match.get('task'):
                qs = qs.filter(task=match['task'])
            if match.get('status'):
                qs = qs.filter(status=match['status'])
            matched = list(qs)
            if not matched:
                result['notes'].append('No treatments matched the update filter')
            for treatment in matched:
                ser = TreatmentSerializer(instance=treatment, data=treatment_fields,
                                          partial=True, context={'request': request})
                ser.is_valid(raise_exception=True)
                ser.save()
                result['treatments_updated'].append(treatment.treatment_id)

        # Operation.
        operation_mode = operation_section.get('mode', 'create')
        operation_fields = self._pick(self.OPERATION_FIELDS, operation_section.get('fields'))
        if operation_mode == 'create' and operation_fields:
            if cohort.op_id:
                result['notes'].append(
                    f"Operation skipped: cohort already has an operation (op_id={cohort.op_id}); use update mode"
                )
            else:
                ser = OperationSerializer(data=operation_fields, context={'request': request})
                ser.is_valid(raise_exception=True)
                operation = ser.save()
                cohort.op_id = operation.op_id
                cohort.save(update_fields=['op_id'])
                assignment.op = operation
                assignment.save(update_fields=['op'])
                result['operation_created'] = operation.op_id
        elif operation_mode == 'update' and operation_fields:
            if not cohort.op_id:
                result['notes'].append('Operation skipped: cohort has no operation linked')
            else:
                operation = self._operation_for_cohort(cohort)
                if operation is None:
                    result['notes'].append(
                        f"Operation skipped: operation (op_id={cohort.op_id}) not found"
                    )
                else:
                    ser = OperationSerializer(instance=operation, data=operation_fields,
                                              partial=True, context={'request': request})
                    ser.is_valid(raise_exception=True)
                    ser.save()
                    result['operation_updated'] = operation.op_id

        return result


class BulkUpdateProposalListView(_BulkUpdateBase):
    """Paginated, searchable proposal list for the bulk-update checkbox UI."""

    def get(self, request):
        qs = Proposal.objects.exclude(processing_status__in=[
            Proposal.PROCESSING_STATUS_DISCARDED,
            Proposal.PROCESSING_STATUS_TEMP,
        ]).order_by('-lodgement_number')

        search = (request.query_params.get('search') or '').strip()
        if search:
            q = Q(lodgement_number__icontains=search)
            if search.isdigit():
                q |= Q(id=int(search))
            qs = qs.filter(q)

        paginator = PageNumberPagination()
        try:
            paginator.page_size = int(request.query_params.get('page_size') or 10)
        except (TypeError, ValueError):
            paginator.page_size = 10
        page = paginator.paginate_queryset(qs, request)
        data = [{'id': p.id, 'lodgement_number': p.lodgement_number} for p in page]
        return paginator.get_paginated_response(data)


class BulkUpdatePolygonListView(_BulkUpdateBase):
    """Paginated polygon list for selected proposal(s), plus an all-ids mode."""

    def get(self, request):
        proposal_ids = self._parse_proposal_ids(request.query_params.get('proposal_ids'))
        if not proposal_ids:
            return Response({'error': 'proposal_ids is required'},
                            status=status.HTTP_400_BAD_REQUEST)

        qs = Polygon.objects.filter(proposal_id__in=proposal_ids).order_by('polygon_id')

        if request.query_params.get('all_ids') == '1':
            ids = [int(x) for x in qs.values_list('polygon_id', flat=True)]
            return Response({'count': len(ids), 'polygon_ids': ids})

        paginator = PageNumberPagination()
        try:
            paginator.page_size = int(request.query_params.get('page_size') or 20)
        except (TypeError, ValueError):
            paginator.page_size = 20
        page = paginator.paginate_queryset(qs, request)
        data = [{'polygon_id': p.polygon_id, 'name': p.name} for p in page]
        return paginator.get_paginated_response(data)

    def _parse_proposal_ids(self, raw):
        if not raw:
            return None
        if isinstance(raw, (list, tuple)):
            tokens = [str(x).strip() for x in raw]
        else:
            tokens = raw.replace(',', ' ').split()
        ids = []
        for token in tokens:
            if not token:
                continue
            if not token.isdigit():
                return None
            ids.append(int(token))
        return ids or None


class BulkUpdateExportTemplateView(_BulkUpdateBase):
    """Export an .xlsx template — one row per resolved polygon, prefilled with
    the current cohort / operation form data plus numbered treatment columns."""

    PREFIX_HEADERS = [
        'polygon_id', 'polygon_name', 'compartment', 'proposal_id', 'cohort_id',
        'obj_code', 'species', 'regen_method', 'site_quality',
        'target_ba_m2ha', 'resid_ba_m2ha', 'target_spha', 'resid_spha',
        'pct_area', 'year_last_cut', 'op_date', 'regen_date', 'regen_date2',
        'regen_done', 'complete_date', 'herbicide_app_spec', 'vrp',
        'vrp_tot_area', 'stand', 'extra_info', 'treatments', 'comments',
        'operation_fea_id', 'operation_das_id', 'operation_plan_release',
    ]

    TREATMENT_FIELDS = [
        'task', 'status', 'plan_yr', 'plan_mth', 'pct_area',
        'complete_date', 'organisation', 'results', 'reference',
    ]

    def post(self, request):
        import io
        from openpyxl import Workbook
        from openpyxl.styles import Font

        ids, err = self._parse_polygon_ids(request.data.get('polygon_ids'))
        if err:
            return Response({'error': err}, status=status.HTTP_400_BAD_REQUEST)

        require_proposal = request.data.get('require_proposal', True)
        polygons, _excluded = self._resolve(ids, require_proposal=require_proposal)

        # Prefetch treatments per cohort and size the numbered treatment columns
        # to the largest count across the selection (always at least one slot).
        for item in polygons:
            item['treatments'] = list(
                Treatment.objects.filter(cohort=item['cohort']).order_by('treatment_id')
            )
        self._n_treatments = max(1, max((len(i['treatments']) for i in polygons), default=1))

        headers = self.PREFIX_HEADERS + self._treatment_headers(self._n_treatments)

        wb = Workbook()
        ws = wb.active
        ws.title = 'Bulk Update'
        ws.append(headers)
        for item in polygons:
            ws.append(self._row_for(item))

        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

        data_end = max(2, len(polygons) + 1)
        lookups = self._fetch_lookups()
        self._build_guide_sheet(wb, lookups)
        lists = self._build_lists_sheet(wb, lookups)
        self._add_dropdowns(ws, lists, data_end)

        buf = io.BytesIO()
        wb.save(buf)
        data = buf.getvalue()
        buf.close()

        response = HttpResponse(
            data,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        filename = f"bulk_update_template_{timezone.now().strftime('%Y%m%dT%H%M')}.xlsx"
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response

    def _fetch_lookups(self):
        """Fetch enumerated dropdown values (code, label) from lookup tables."""
        return {
            'objectives': list(
                ObjectiveLkp.objects.values_list('obj_code', 'description').order_by('obj_code')
            ),
            'species': list(
                SpeciesApiLkp.objects.values_list('species', 'short_description').order_by('species')
            ),
            'regen_methods': list(
                RegenerationMethodsLkp.objects.values_list('regen_method', 'description').order_by('regen_method')
            ),
            'tasks': list(
                TaskLkp.objects.values_list('task', 'task_name').order_by('task')
            ),
            'statuses': [
                ('P', 'Planned'),
                ('D', 'Completed'),
                ('C', 'Cancelled'),
                ('F', 'Failed'),
                ('W', 'Written Off'),
                ('X', 'Not Required'),
            ],
        }

    def _build_lists_sheet(self, wb, lookups):
        """Hidden sheet holding validation lists (avoids the 255-char inline
        limit for long lookup lists). Returns {name: formula1 range}."""
        from openpyxl.utils import get_column_letter

        lists_ws = wb.create_sheet('Lists')
        lists_ws.sheet_state = 'hidden'

        ranges = {}
        col = 1
        for name in ('objectives', 'species', 'regen_methods', 'tasks', 'statuses'):
            codes = [code.strip() for code, _label in lookups[name] if code]
            if not codes:
                codes = [' ']
            for row_idx, value in enumerate(codes, start=1):
                lists_ws.cell(row=row_idx, column=col, value=value)
            letter = get_column_letter(col)
            ranges[name] = f'Lists!${letter}$1:${letter}${len(codes)}'
            col += 1
        return ranges

    def _add_dropdowns(self, ws, lists, data_end):
        """Add in-cell list dropdowns to the enumerated columns."""
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation

        def apply(col_idx, formula1):
            letter = get_column_letter(col_idx + 1)
            dv = DataValidation(type='list', formula1=formula1, allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f'{letter}2:{letter}{data_end}')

        # Cohort enumerated fields
        apply(5, lists['objectives'])      # obj_code
        apply(6, lists['species'])         # species
        apply(7, lists['regen_methods'])   # regen_method
        apply(18, '"Yes,No"')              # regen_done
        apply(24, '"Yes,No"')              # extra_info
        apply(25, '"Yes,No"')              # treatments

        # Numbered treatment slots (id + 9 fields each)
        for i in range(self._n_treatments):
            base = 30 + i * 10
            apply(base + 1, lists['tasks'])      # treatment_i_task
            apply(base + 2, lists['statuses'])   # treatment_i_status

    def _build_guide_sheet(self, wb, lookups):
        from openpyxl.styles import Font, PatternFill

        guide = wb.create_sheet('Guide')
        bold = Font(bold=True)
        title_font = Font(bold=True, size=14)
        header_fill = PatternFill(start_color='DDEBF7', end_color='DDEBF7', fill_type='solid')

        guide.cell(row=1, column=1, value='Bulk Update Template — How to Use').font = title_font

        lines = [
            'The "Bulk Update" sheet has one row per polygon, prefilled with current data.',
            '',
            'Rules',
            '- Leave a cell BLANK to keep its current value (only filled cells are updated).',
            '- Cells with a dropdown (▼) accept only the listed values — pick from the list.',
            '- Do not edit reference columns: Polygon ID, Name, Compartment, Proposal, Cohort ID, or treatment_N_id.',
            '- Dates use YYYY-MM-DD HH:MM. Yes/No fields use Yes or No.',
            '- Treatments are numbered (treatment_1_*, treatment_2_*, ...). Each block holds one treatment.',
            '  A block with a treatment_N_id updates that existing treatment; a block with no id but a',
            '  Task adds a new treatment. Leave a block blank to make no treatment changes.',
            '- Upload the completed file via the "Excel Upload" tab on the Bulk Update page.',
            '',
            'Dropdown value reference',
        ]
        row = 3
        for line in lines:
            guide.cell(row=row, column=1, value=line)
            row += 1

        refs = [
            ('Objective (obj_code)', lookups['objectives']),
            ('Species', lookups['species']),
            ('Regeneration Method (regen_method)', lookups['regen_methods']),
            ('Task (treatment_*_task)', lookups['tasks']),
            ('Treatment Status (treatment_*_status)', lookups['statuses']),
            ('Yes/No', [('Yes', ''), ('No', '')]),
        ]
        for title, values in refs:
            c1 = guide.cell(row=row, column=1, value=title)
            c1.font = bold
            c1.fill = header_fill
            c2 = guide.cell(row=row, column=2, value='Description')
            c2.font = bold
            c2.fill = header_fill
            row += 1
            for code, label in values:
                guide.cell(row=row, column=1, value=code)
                guide.cell(row=row, column=2, value=label or '')
                row += 1
            row += 1

        guide.column_dimensions['A'].width = 45
        guide.column_dimensions['B'].width = 70

    def _treatment_headers(self, n):
        headers = []
        for i in range(1, n + 1):
            headers.append(f'treatment_{i}_id')
            for field in self.TREATMENT_FIELDS:
                headers.append(f'treatment_{i}_{field}')
        return headers

    def _fmt_dt(self, val):
        if not val:
            return ''
        return val.strftime('%Y-%m-%d %H:%M')

    def _fmt_bool(self, val):
        if val is None:
            return ''
        return 'Yes' if val else 'No'

    @staticmethod
    def _str(value):
        return value.strip() if isinstance(value, str) else value

    def _treatment_values(self, treatment):
        if treatment is None:
            return [''] + [''] * len(self.TREATMENT_FIELDS)
        return [
            treatment.treatment_id,
            self._str(treatment.task_id or ''),
            self._str(treatment.status or ''),
            treatment.plan_yr,
            treatment.plan_mth,
            treatment.pct_area,
            self._fmt_dt(treatment.complete_date),
            self._str(treatment.organisation or ''),
            self._str(treatment.results or ''),
            self._str(treatment.reference or ''),
        ]

    def _row_for(self, item):
        poly, cohort = item['polygon'], item['cohort']
        op = self._operation_for_cohort(cohort)
        treatments = item.get('treatments') or []

        row = [
            poly.polygon_id,
            self._str(poly.name or ''),
            self._str(poly.compartment.compartment if poly.compartment else ''),
            poly.proposal_id or '',
            cohort.cohort_id,
            self._str(cohort.obj_code or ''),
            self._str(cohort.species or ''),
            self._str(cohort.regen_method_id or ''),
            self._str(cohort.site_quality or ''),
            cohort.target_ba_m2ha,
            cohort.resid_ba_m2ha,
            cohort.target_spha,
            cohort.resid_spha,
            cohort.pct_area,
            cohort.year_last_cut,
            self._fmt_dt(cohort.op_date),
            self._fmt_dt(cohort.regen_date),
            self._fmt_dt(cohort.regen_date2),
            self._fmt_bool(cohort.regen_done),
            self._fmt_dt(cohort.complete_date),
            self._str(cohort.herbicide_app_spec or ''),
            self._str(cohort.vrp_id or ''),
            cohort.vrp_tot_area,
            self._str(cohort.stand or ''),
            self._fmt_bool(cohort.extra_info),
            self._fmt_bool(cohort.treatments),
            self._str(cohort.comments or ''),
            self._str(op.fea_id if op else ''),
            op.das_id if op else '',
            self._str(op.plan_release if op else ''),
        ]

        for i in range(self._n_treatments):
            treatment = treatments[i] if i < len(treatments) else None
            row.extend(self._treatment_values(treatment))

        return row


class BulkUpdateUploadError(Exception):
    """Raised when applying an uploaded template hits validation errors.

    Carries a list of {row, column, message} dicts so the client can report
    the exact cell location of each problem.
    """

    def __init__(self, errors):
        self.errors = errors
        super().__init__('; '.join(f"{e['row']}:{e['column']} {e['message']}" for e in errors))


class BulkUpdateUploadView(_BulkUpdateBase):
    """Parse and apply an uploaded bulk-update template, reporting per-cell
    validation errors with their Excel row/column location."""

    parser_classes = [MultiPartParser, FormParser, JSONParser]

    COHORT_FIELD_TYPES = {
        'obj_code': ('lookup', 'objectives'),
        'species': ('lookup', 'species'),
        'regen_method': ('lookup', 'regen_methods'),
        'site_quality': ('text', None),
        'target_ba_m2ha': ('float', None),
        'resid_ba_m2ha': ('float', None),
        'target_spha': ('int', None),
        'resid_spha': ('float', None),
        'pct_area': ('pct', None),
        'year_last_cut': ('int', None),
        'op_date': ('date', None),
        'regen_date': ('date', None),
        'regen_date2': ('date', None),
        'regen_done': ('bool', None),
        'complete_date': ('date', None),
        'herbicide_app_spec': ('text', None),
        'vrp': ('int', None),
        'vrp_tot_area': ('float', None),
        'stand': ('text', None),
        'extra_info': ('bool', None),
        'treatments': ('bool', None),
        'comments': ('text', None),
    }

    OPERATION_FIELD_TYPES = {
        'operation_fea_id': ('fea_id', 'text'),
        'operation_das_id': ('das_id', 'int'),
        'operation_plan_release': ('plan_release', 'text'),
    }

    TREATMENT_FIELDS = [
        ('task', ('lookup', 'tasks')),
        ('status', ('lookup', 'statuses')),
        ('plan_yr', ('int', None)),
        ('plan_mth', ('month', None)),
        ('pct_area', ('pct', None)),
        ('complete_date', ('date', None)),
        ('organisation', ('text', None)),
        ('results', ('text', None)),
        ('reference', ('text', None)),
    ]

    def post(self, request):
        import io
        import re
        from openpyxl import load_workbook

        uploaded = request.FILES.get('file')
        if uploaded is None:
            return Response({'error': 'No file uploaded'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            wb = load_workbook(io.BytesIO(uploaded.read()), data_only=True)
        except Exception:
            return Response(
                {'error': 'Could not read the uploaded file. Upload a valid .xlsx template.'},
                status=status.HTTP_400_BAD_REQUEST,
            )

        ws = wb['Bulk Update'] if 'Bulk Update' in wb.sheetnames else wb.active

        headers = {}
        for col_idx, cell in enumerate(ws[1], start=1):
            if cell.value is not None and str(cell.value).strip():
                headers[str(cell.value).strip()] = col_idx

        max_treatments = 0
        for h in headers:
            m = re.match(r'^treatment_(\d+)_task$', h)
            if m:
                max_treatments = max(max_treatments, int(m.group(1)))
        self._max_treatments = max_treatments

        lookups = self._fetch_lookups()
        lookup_maps = {k: {code.strip(): code for code, _label in v if code} for k, v in lookups.items()}

        errors = []
        parsed_rows = []
        for row_idx in range(2, ws.max_row + 1):
            values = [ws.cell(row=row_idx, column=c).value for c in range(1, ws.max_column + 1)]
            if all(v is None or str(v).strip() == '' for v in values):
                continue

            row_errors, parsed = self._parse_row(ws, row_idx, headers, lookup_maps)
            errors.extend(row_errors)
            if row_errors:
                continue

            polygons, excluded = self._resolve(
                [parsed['polygon_id']], require_proposal=False
            )
            if not polygons:
                errors.append({
                    'row': row_idx,
                    'column': 'polygon_id',
                    'message': excluded[0]['reason'] if excluded else 'Polygon not found',
                })
                continue
            bad_treatment = False
            for block in parsed['treatments']:
                if 'id' in block and not Treatment.objects.filter(
                        treatment_id=block['id'], cohort=polygons[0]['cohort']).exists():
                    errors.append({
                        'row': row_idx,
                        'column': block.get('_id_col', 'treatment_id'),
                        'message': f"Treatment {block['id']} does not belong to cohort {polygons[0]['cohort'].cohort_id}",
                    })
                    bad_treatment = True
            if bad_treatment:
                continue
            parsed['item'] = polygons[0]
            parsed['row'] = row_idx
            parsed_rows.append(parsed)

        if errors:
            return Response(
                {'error': f'{len(errors)} validation error(s) found', 'errors': errors},
                status=status.HTTP_400_BAD_REQUEST,
            )

        results = []
        try:
            with transaction.atomic():
                for parsed in parsed_rows:
                    results.append(self._apply_row(request, parsed))
        except BulkUpdateUploadError as exc:
            return Response(
                {'error': f'{len(exc.errors)} validation error(s) found', 'errors': exc.errors},
                status=status.HTTP_400_BAD_REQUEST,
            )
        except serializers.ValidationError as exc:
            return Response({'error': 'Validation failed', 'details': exc.detail},
                            status=status.HTTP_400_BAD_REQUEST)
        except Exception:
            logger.exception('Excel upload failed and was rolled back')
            return Response(
                {'error': 'Upload failed and was rolled back. No changes were applied.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

        return Response({'processed': len(results), 'results': results})

    def _fetch_lookups(self):
        """Fetch enumerated dropdown values (code, label) from lookup tables."""
        return {
            'objectives': list(
                ObjectiveLkp.objects.values_list('obj_code', 'description').order_by('obj_code')
            ),
            'species': list(
                SpeciesApiLkp.objects.values_list('species', 'short_description').order_by('species')
            ),
            'regen_methods': list(
                RegenerationMethodsLkp.objects.values_list('regen_method', 'description').order_by('regen_method')
            ),
            'tasks': list(
                TaskLkp.objects.values_list('task', 'task_name').order_by('task')
            ),
            'statuses': [
                ('P', 'Planned'),
                ('D', 'Completed'),
                ('C', 'Cancelled'),
                ('F', 'Failed'),
                ('W', 'Written Off'),
                ('X', 'Not Required'),
            ],
        }

    def _parse_row(self, ws, row_idx, headers, lookup_maps):
        errors = []
        parsed = {'polygon_id': None, 'cohort_fields': {}, 'operation_fields': {}, 'treatments': []}

        def read(name):
            col = headers.get(name)
            return ws.cell(row=row_idx, column=col).value if col else None

        pid, err = self._clean(read('polygon_id'), 'int', None)
        if err:
            errors.append(self._err(row_idx, 'polygon_id', err))
        elif pid is None:
            errors.append(self._err(row_idx, 'polygon_id', 'Polygon ID is required'))
        else:
            parsed['polygon_id'] = pid

        for name, (type_, lookup_name) in self.COHORT_FIELD_TYPES.items():
            cleaned, err = self._clean(read(name), type_,
                                       lookup_maps.get(lookup_name) if lookup_name else None)
            if err:
                errors.append(self._err(row_idx, name, err))
            elif cleaned is not None:
                parsed['cohort_fields'][name] = cleaned

        for col_name, (field_name, type_) in self.OPERATION_FIELD_TYPES.items():
            cleaned, err = self._clean(read(col_name), type_, None)
            if err:
                errors.append(self._err(row_idx, col_name, err))
            elif cleaned is not None:
                parsed['operation_fields'][field_name] = cleaned

        for i in range(1, self._max_treatments + 1):
            block = {'_num': i}
            tid, err = self._clean(read(f'treatment_{i}_id'), 'int', None)
            if err:
                errors.append(self._err(row_idx, f'treatment_{i}_id', err))
            elif tid is not None:
                block['id'] = tid
                block['_id_col'] = f'treatment_{i}_id'
            for field_name, (type_, lookup_name) in self.TREATMENT_FIELDS:
                hname = f'treatment_{i}_{field_name}'
                cleaned, err = self._clean(read(hname), type_,
                                           lookup_maps.get(lookup_name) if lookup_name else None)
                if err:
                    errors.append(self._err(row_idx, hname, err))
                elif cleaned is not None:
                    block[field_name] = cleaned
            has_fields = any(k not in ('id', '_id_col', '_num') for k in block)
            if block.get('id') is not None:
                if has_fields:
                    parsed['treatments'].append(block)
            elif has_fields:
                if 'task' not in block:
                    errors.append(self._err(row_idx, f'treatment_{i}_task',
                                            'Treatment block has values but no Task'))
                else:
                    parsed['treatments'].append(block)

        return errors, parsed

    def _clean(self, value, type_, lookup_map):
        """Return (cleaned_value, error_message_or_None)."""
        if value is None or (isinstance(value, str) and value.strip() == ''):
            return None, None
        if isinstance(value, str):
            value = value.strip()

        if type_ == 'text':
            return value, None
        if type_ == 'lookup':
            if value not in lookup_map:
                return None, f"'{value}' is not a valid option"
            return lookup_map[value], None
        if type_ == 'bool':
            if isinstance(value, bool):
                return value, None
            s = str(value).strip().lower()
            if s in ('yes', 'true', '1'):
                return True, None
            if s in ('no', 'false', '0'):
                return False, None
            return None, f"'{value}' must be Yes or No"
        if type_ == 'int':
            if isinstance(value, bool):
                return None, 'must be a whole number'
            if isinstance(value, float) and not value.is_integer():
                return None, 'must be a whole number'
            try:
                return int(value), None
            except (TypeError, ValueError):
                return None, f"'{value}' is not a whole number"
        if type_ == 'float':
            if isinstance(value, bool):
                return None, 'must be a number'
            try:
                return float(value), None
            except (TypeError, ValueError):
                return None, f"'{value}' is not a number"
        if type_ == 'pct':
            if isinstance(value, bool):
                return None, 'must be a whole number between 0 and 100'
            if isinstance(value, float) and not value.is_integer():
                return None, 'must be a whole number between 0 and 100'
            try:
                n = int(value)
            except (TypeError, ValueError):
                return None, f"'{value}' is not a whole number"
            if n < 0 or n > 100:
                return None, 'must be between 0 and 100'
            return n, None
        if type_ == 'month':
            if isinstance(value, bool):
                return None, 'must be a number between 1 and 12'
            if isinstance(value, float) and not value.is_integer():
                return None, 'must be a whole number'
            try:
                n = int(value)
            except (TypeError, ValueError):
                return None, f"'{value}' is not a number"
            if n < 1 or n > 12:
                return None, 'must be between 1 and 12'
            return n, None
        if type_ == 'date':
            if isinstance(value, (datetime, date)):
                return value, None
            s = str(value).strip()
            for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y-%m-%d'):
                try:
                    return datetime.strptime(s, fmt), None
                except ValueError:
                    continue
            return None, f"'{s}' is not a valid date (use YYYY-MM-DD or YYYY-MM-DD HH:MM)"

        return value, None

    def _err(self, row_idx, column, message):
        return {'row': row_idx, 'column': column, 'message': message}

    def _loc_errors(self, row_idx, prefix, detail):
        """Convert DRF serializer error detail to row/column location errors."""
        errors = []
        for field, msgs in (detail or {}).items():
            if isinstance(msgs, (list, tuple)):
                msg = '; '.join(str(m) for m in msgs)
            else:
                msg = str(msgs)
            col = f'{prefix}_{field}' if prefix else field
            errors.append({'row': row_idx, 'column': col, 'message': msg})
        return errors

    @staticmethod
    def _display_value(value):
        """Normalise a value for old/new comparison and display."""
        if value is None:
            return None
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if hasattr(value, '_meta'):  # Django model instance (FK)
            return value.pk
        if isinstance(value, (int, float, str, bool)):
            return value
        return str(value)

    def _field_changes(self, obj, fields):
        """Diff the given {field: new_value} against obj's current values."""
        changes = []
        for field, new_value in fields.items():
            old = self._display_value(getattr(obj, field, None))
            new = self._display_value(new_value)
            if old != new:
                changes.append({'field': field, 'old': old, 'new': new})
        return changes

    def _apply_row(self, request, parsed):
        item = parsed['item']
        poly, cohort, assignment = item['polygon'], item['cohort'], item['assignment']
        row_idx = parsed.get('row')
        result = {
            'polygon_id': poly.polygon_id,
            'cohort_id': cohort.cohort_id,
            'cohort_updated': False,
            'cohort_changes': [],
            'treatments_created': [],
            'treatments_updated': [],
            'treatment_changes': [],
            'operation_created': None,
            'operation_updated': None,
            'operation_changes': [],
        }

        if parsed['cohort_fields']:
            ser = CohortSerializer(instance=cohort, data=parsed['cohort_fields'],
                                   partial=True, context={'request': request})
            if not ser.is_valid():
                raise BulkUpdateUploadError(self._loc_errors(row_idx, '', ser.errors))
            result['cohort_changes'] = self._field_changes(cohort, parsed['cohort_fields'])
            ser.save()
            result['cohort_updated'] = bool(result['cohort_changes'])

        for block in parsed['treatments']:
            tid = block.get('id')
            prefix = f"treatment_{block.get('_num', '?')}"
            fields = {k: v for k, v in block.items() if k not in ('id', '_id_col', '_num')}
            if tid is not None:
                try:
                    treatment = Treatment.objects.get(treatment_id=tid, cohort=cohort)
                except Treatment.DoesNotExist:
                    raise BulkUpdateUploadError([{
                        'row': row_idx,
                        'column': block.get('_id_col', 'treatment_id'),
                        'message': f"Treatment {tid} does not belong to cohort {cohort.cohort_id}",
                    }])
                ser = TreatmentSerializer(instance=treatment, data=fields,
                                          partial=True, context={'request': request})
                if not ser.is_valid():
                    raise BulkUpdateUploadError(self._loc_errors(row_idx, prefix, ser.errors))
                changes = self._field_changes(treatment, fields)
                treatment = ser.save()
                result['treatment_changes'].append({'treatment_id': tid, 'changes': changes})
                result['treatments_updated'].append(treatment.treatment_id)
            else:
                ser = TreatmentSerializer(data={**fields, 'cohort': cohort.cohort_id},
                                          context={'request': request})
                if not ser.is_valid():
                    raise BulkUpdateUploadError(self._loc_errors(row_idx, prefix, ser.errors))
                treatment = ser.save()
                result['treatments_created'].append(treatment.treatment_id)

        op_fields = parsed['operation_fields']
        if op_fields:
            if cohort.op_id:
                op = self._operation_for_cohort(cohort)
                if op is not None:
                    ser = OperationSerializer(instance=op, data=op_fields,
                                              partial=True, context={'request': request})
                    if not ser.is_valid():
                        raise BulkUpdateUploadError(
                            self._loc_errors(row_idx, 'operation', ser.errors)
                        )
                    result['operation_changes'] = self._field_changes(op, op_fields)
                    ser.save()
                    result['operation_updated'] = op.op_id
            else:
                ser = OperationSerializer(data=op_fields, context={'request': request})
                if not ser.is_valid():
                    raise BulkUpdateUploadError(
                        self._loc_errors(row_idx, 'operation', ser.errors)
                    )
                op = ser.save()
                cohort.op_id = op.op_id
                cohort.save(update_fields=['op_id'])
                assignment.op = op
                assignment.save(update_fields=['op'])
                result['operation_created'] = op.op_id

        return result


class ERDiagramView(APIView):
    """Return model/field/relationship metadata for the ER diagram viewer.

    Models are grouped by the module that defines them so the front-end can
    present "Main", "Proposals", "Forest Blocks" and "Lookups" separately
    even though Main and Proposals both register under the 'silrec' app label.
    """

    permission_classes = [IsAuthenticated & (IsOperator | IsReviewer | IsSilrecAdmin)]

    MODULE_GROUPS = [
        ('main', 'Main', 'silrec.components.main.models'),
        ('proposals', 'Proposals', 'silrec.components.proposals.models'),
        ('forest_blocks', 'Forest Blocks', 'silrec.components.forest_blocks.models'),
        ('lookups', 'Lookups', 'silrec.components.lookups.models'),
    ]

    def get(self, request):
        from collections import defaultdict
        from django.apps import apps

        by_module = defaultdict(list)
        for model in apps.get_models():
            if model._meta.abstract or model._meta.proxy:
                continue
            by_module[model.__module__].append(model)

        groups = []
        for key, display, module in self.MODULE_GROUPS:
            models = sorted(by_module.get(module, []), key=lambda m: m.__name__.lower())
            groups.append({
                'key': key,
                'name': display,
                'models': [self._model_info(m) for m in models],
            })
        return Response({'groups': groups})

    def _model_info(self, model):
        fields = []
        for field in model._meta.fields:
            info = {
                'name': field.name,
                'type': field.__class__.__name__,
                'pk': bool(field.primary_key),
            }
            if field.is_relation and field.related_model is not None:
                info['related_model'] = field.related_model.__name__
                info['relation'] = 'o2o' if field.one_to_one else 'fk'
            fields.append(info)
        for field in model._meta.many_to_many:
            info = {
                'name': field.name,
                'type': 'ManyToManyField',
                'pk': False,
                'relation': 'm2m',
            }
            if field.related_model is not None:
                info['related_model'] = field.related_model.__name__
            fields.append(info)
        return {
            'name': model.__name__,
            'verbose_name': str(model._meta.verbose_name),
            'db_table': model._meta.db_table,
            'fields': fields,
        }
